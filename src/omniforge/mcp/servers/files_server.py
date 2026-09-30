"""MCP server exposing document reading and file inspection tools."""
import json
import logging
import mimetypes
import os
from pathlib import Path
import time
from typing import Any

logger = logging.getLogger(__name__)


def _format_size(size_bytes: int) -> str:
    """Format bytes into human-readable string."""
    for unit in ["B", "KB", "MB", "GB"]:
        if abs(size_bytes) < 1024.0:
            return f"{size_bytes:3.1f} {unit}"
        size_bytes /= 1024.0
    return f"{size_bytes:.1f} TB"


async def list_files(
    directory: str = ".",
    pattern: str = "*",
    recursive: bool = True,
    max_results: int = 100,
) -> str:
    """List and search files matching a glob pattern within a directory.

    Args:
        directory: Root directory to search within.
        pattern: Glob pattern to filter filenames (e.g. '*.py', '*.pdf', '*report*').
        recursive: Whether to search recursively in subdirectories.
        max_results: Maximum number of files to return.

    Returns:
        JSON string containing matched files and their metadata.
    """
    try:
        base_dir = Path(directory).resolve()
        if not base_dir.exists():
            return json.dumps({"error": f"Directory not found: {directory}"})
        if not base_dir.is_dir():
            return json.dumps({"error": f"Path is not a directory: {directory}"})

        iterator = base_dir.rglob(pattern) if recursive else base_dir.glob(pattern)
        results: list[dict[str, Any]] = []

        for p in iterator:
            if p.is_file():
                if len(results) >= max_results:
                    break
                try:
                    stat = p.stat()
                    results.append(
                        {
                            "name": p.name,
                            "relative_path": str(p.relative_to(base_dir)),
                            "extension": p.suffix.lower(),
                            "size_bytes": stat.st_size,
                            "size_human": _format_size(stat.st_size),
                            "modified": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_mtime)),
                        }
                    )
                except (OSError, PermissionError):
                    continue

        return json.dumps(
            {
                "directory": str(base_dir),
                "pattern": pattern,
                "recursive": recursive,
                "total_matched": len(results),
                "files": results,
            },
            indent=2,
        )
    except Exception as e:
        logger.error("list_files failed: %s", e)
        return json.dumps({"error": str(e)})


async def read_document(path: str, max_chars: int = 30000) -> str:
    """Read and extract text from documents (PDF, DOCX, TXT, CSV, Excel, Markdown, JSON).

    Args:
        path: Path to the target document.
        max_chars: Maximum characters to return.

    Returns:
        Extracted textual content or error message.
    """
    try:
        file_path = Path(path).resolve()
        if not file_path.exists():
            return f"Error: File not found: {path}"
        if not file_path.is_file():
            return f"Error: Path is not a file: {path}"

        suffix = file_path.suffix.lower()
        extracted_text = ""

        # PDF handling via pypdf
        if suffix == ".pdf":
            try:
                import pypdf

                reader = pypdf.PdfReader(str(file_path))
                pages_text = []
                for i, page in enumerate(reader.pages):
                    page_content = page.extract_text() or ""
                    pages_text.append(f"--- Page {i + 1} ---\n{page_content}")
                extracted_text = "\n\n".join(pages_text)
            except ImportError:
                return "Error: pypdf library is not installed to read PDF files."

        # Word document handling via python-docx
        elif suffix in [".docx", ".doc"]:
            try:
                import docx

                doc = docx.Document(str(file_path))
                paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
                # Also include table cells
                table_texts = []
                for t in doc.tables:
                    for row in t.rows:
                        row_vals = [c.text.strip() for c in row.cells if c.text.strip()]
                        if row_vals:
                            table_texts.append(" | ".join(row_vals))
                extracted_text = "\n".join(paragraphs)
                if table_texts:
                    extracted_text += "\n\n--- Tables ---\n" + "\n".join(table_texts)
            except ImportError:
                return "Error: python-docx library is not installed to read DOCX files."

        # Tabular data via pandas or openpyxl
        elif suffix in [".xlsx", ".xls"]:
            try:
                import pandas as pd

                sheets = pd.read_excel(str(file_path), sheet_name=None)
                sheet_summaries = []
                for sheet_name, df in sheets.items():
                    sheet_summaries.append(
                        f"### Sheet: {sheet_name} (Shape: {df.shape[0]} rows, {df.shape[1]} cols)\n"
                        f"{df.head(10).to_markdown(index=False)}"
                    )
                extracted_text = "\n\n".join(sheet_summaries)
            except ImportError:
                return "Error: pandas/openpyxl is not installed to read Excel files."

        elif suffix == ".csv":
            try:
                import pandas as pd

                df = pd.read_csv(str(file_path))
                extracted_text = (
                    f"### CSV File Summary (Shape: {df.shape[0]} rows, {df.shape[1]} cols)\n"
                    f"{df.head(20).to_markdown(index=False)}"
                )
            except Exception:
                extracted_text = file_path.read_text(encoding="utf-8", errors="replace")

        # Standard plain text / markdown / code / json
        else:
            try:
                extracted_text = file_path.read_text(encoding="utf-8")
            except UnicodeDecodeError:
                extracted_text = file_path.read_text(encoding="latin-1", errors="replace")

        if not extracted_text.strip():
            return f"Document '{path}' is empty or contains no extractable text."

        if len(extracted_text) > max_chars:
            extracted_text = (
                extracted_text[:max_chars]
                + f"\n\n... [Content truncated at {max_chars} of {len(extracted_text)} total characters]"
            )

        return extracted_text
    except Exception as e:
        logger.error("read_document failed for %s: %s", path, e)
        return f"Error reading document '{path}': {e}"


async def get_file_info(path: str) -> str:
    """Retrieve detailed metadata and filesystem statistics for a file or directory.

    Args:
        path: Path to the target file or directory.

    Returns:
        JSON string containing size, timestamps, mime type, and properties.
    """
    try:
        p = Path(path).resolve()
        if not p.exists():
            return json.dumps({"exists": False, "path": str(p), "error": "Path does not exist"})

        stat = p.stat()
        mime_type, _ = mimetypes.guess_type(str(p))

        line_count = None
        if p.is_file() and p.suffix.lower() in [
            ".py", ".txt", ".md", ".json", ".yaml", ".yml", ".csv", ".toml", ".ini", ".html", ".css", ".js"
        ]:
            try:
                with p.open("r", encoding="utf-8", errors="ignore") as f:
                    line_count = sum(1 for _ in f)
            except Exception:
                pass

        info = {
            "exists": True,
            "name": p.name,
            "path": str(p),
            "is_file": p.is_file(),
            "is_dir": p.is_dir(),
            "size_bytes": stat.st_size if p.is_file() else 0,
            "size_human": _format_size(stat.st_size) if p.is_file() else "N/A",
            "extension": p.suffix.lower() if p.is_file() else "",
            "mime_type": mime_type or "unknown",
            "created": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_ctime)),
            "modified": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_mtime)),
            "line_count": line_count,
        }
        return json.dumps(info, indent=2)
    except Exception as e:
        logger.error("get_file_info failed for %s: %s", path, e)
        return json.dumps({"error": str(e)})


try:
    from mcp.server.fastmcp import FastMCP

    mcp = FastMCP("OmniForge Files")
    mcp.tool()(list_files)
    mcp.tool()(read_document)
    mcp.tool()(get_file_info)

    if __name__ == "__main__":
        mcp.run()
except ImportError:
    mcp = None  # MCP library not installed or unavailable in current environment
