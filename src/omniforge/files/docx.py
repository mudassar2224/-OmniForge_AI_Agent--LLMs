"""DOCX document text and metadata extraction using python-docx."""

import logging
from pathlib import Path
from typing import Any
from pydantic import BaseModel, Field

logger = logging.getLogger(__name__)


class DocxMetadata(BaseModel):
    """Metadata extracted from a DOCX document."""

    paragraphs: int = Field(default=0, description="Total number of paragraphs")
    tables: int = Field(default=0, description="Total number of tables")
    title: str | None = Field(default=None, description="Document title")
    author: str | None = Field(default=None, description="Document author")
    subject: str | None = Field(default=None, description="Document subject")
    created: str | None = Field(default=None, description="Creation timestamp")
    modified: str | None = Field(default=None, description="Modification timestamp")
    category: str | None = Field(default=None, description="Document category")


def extract_docx_text(file_path: str | Path) -> str:
    """Extract all text from a DOCX file including paragraphs and tables.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Extracted text formatted with paragraphs and table cells.
    """
    path = Path(file_path)
    if not path.exists():
        logger.error("DOCX file not found: %s", file_path)
        return f"Error extracting DOCX: File not found: {file_path}"

    try:
        from docx import Document

        doc = Document(str(path))
        text_parts: list[str] = []

        # Extract text from paragraphs
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_parts.append(p.text.strip())

        # Extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    text_parts.append(" | ".join(row_cells))

        return "\n\n".join(text_parts)
    except Exception as e:
        logger.error("DOCX extraction failed for %s: %s", file_path, e)
        return f"Error extracting DOCX: {e}"


def get_docx_metadata(file_path: str | Path) -> dict[str, Any]:
    """Retrieve core metadata and structural counts from a DOCX file.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Dictionary containing title, author, subject, creation date, and counts.
    """
    path = Path(file_path)
    if not path.exists():
        return {"error": f"File not found: {file_path}"}

    try:
        from docx import Document

        doc = Document(str(path))
        props = doc.core_properties

        created_str = props.created.isoformat() if props.created else None
        modified_str = props.modified.isoformat() if props.modified else None

        metadata = DocxMetadata(
            paragraphs=len(doc.paragraphs),
            tables=len(doc.tables),
            title=props.title if props.title else None,
            author=props.author if props.author else None,
            subject=props.subject if props.subject else None,
            created=created_str,
            modified=modified_str,
            category=props.category if props.category else None,
        )
        return metadata.model_dump()
    except Exception as e:
        logger.error("Failed to read DOCX metadata for %s: %s", file_path, e)
        return {"error": str(e)}
